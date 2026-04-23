import os, json
from datetime import datetime
import openpyxl, xlrd

base = 'D:/table/课堂练习 2'

# 找 xls 文件
files = [f for f in os.listdir(base) if f.endswith('.xlsx') or f.endswith('.xls')]
print('找到文件:', files)

result = []

for fn in files:
    fp = os.path.join(base, fn)
    print('处理:', fn, '存在:', os.path.exists(fp))
    if not os.path.exists(fp):
        continue
    tables = []
    try:
        if fn.endswith('.xlsx'):
            wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
            for sh in wb.sheetnames:
                ws = wb[sh]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) if c else '' for c in row]
                    while cells and cells[-1] == '': cells.pop()
                    if cells and any(c.strip() for c in cells): rows.append(' | '.join(cells))
                if rows:
                    tables.append('=== ' + sh + ' ===')
                    tables.extend(rows)
        else:
            wb = xlrd.open_workbook(fp)
            for sh in wb.sheet_names():
                ws = wb.sheet_by_name(sh)
                rows = []
                for ri in range(ws.nrows):
                    cells = [str(ws.cell_value(ri, ci)).strip() for ci in range(ws.ncols)]
                    while cells and cells[-1] == '': cells.pop()
                    if cells and any(c.strip() for c in cells): rows.append(' | '.join(cells))
                if rows:
                    tables.append('=== ' + sh + ' ===')
                    tables.extend(rows)
        txt = '\n'.join(tables)
        print('提取:', len(txt), '字')
        st = os.stat(fp)
        result.append({'name': fn, 'relativePath': fn, 'type': 'excel', 'size': st.st_size, 'lastModified': datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d %H:%M:%S'), 'contentBase64': None, 'textContent': txt or None})
    except Exception as e:
        print('错误:', e)

out = os.path.join(base, '7. 2026 年履职计划.json')
with open(out, 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)
print('保存到:', out)
print('文件数:', len(result))

# 复制 to src
import shutil
shutil.copy(out, os.path.join(base, 'src', '7. 2026 年履职计划.json'))
print('已复制到 src')
